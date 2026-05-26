#!/usr/bin/env python3
"""
Скрипт для конвертации CSV в отформатированный Excel файл
"""
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.utils import get_column_letter

# Чтение CSV файла
print("Чтение CSV файла...")
df = pd.read_csv('Teen_Mental_Health_Dataset.csv')

# Создание Excel файла
excel_filename = 'Teen_Mental_Health_Dataset.xlsx'
print(f"Создание Excel файла: {excel_filename}")

# Сохраняем DataFrame в Excel
df.to_excel(excel_filename, index=False, sheet_name='Data')

# Загрузка workbook для форматирования
wb = load_workbook(excel_filename)
ws = wb.active

# Определение стилей
header_font = Font(bold=True, color="FFFFFF", size=12)
header_alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")

cell_alignment = Alignment(horizontal='left', vertical='center')
thin_border = Border(
    left=Side(style='thin'),
    right=Side(style='thin'),
    top=Side(style='thin'),
    bottom=Side(style='thin')
)

# Форматирование заголовков
for col in range(1, len(df.columns) + 1):
    cell = ws.cell(row=1, column=col)
    cell.font = header_font
    cell.alignment = header_alignment
    cell.fill = header_fill
    cell.border = thin_border

# Форматирование данных
for row in range(2, ws.max_row + 1):
    for col in range(1, ws.max_column + 1):
        cell = ws.cell(row=row, column=col)
        cell.alignment = cell_alignment
        cell.border = thin_border

# Автоподбор ширины столбцов
for col in ws.columns:
    max_length = 0
    column = col[0].column_letter
    
    # Находим максимальную длину значения в столбце
    for cell in col:
        if cell.value:
            try:
                cell_length = len(str(cell.value))
                if cell_length > max_length:
                    max_length = cell_length
            except:
                pass
    
    # Устанавливаем ширину столбца с небольшим запасом
    adjusted_width = min(max_length + 2, 50)  # Максимум 50 символов
    ws.column_dimensions[column].width = adjusted_width

# Фиксация первой строки (заголовки)
ws.freeze_panes = 'A2'

# Сохранение файла
wb.save(excel_filename)
print(f"Excel файл успешно создан и отформатирован: {excel_filename}")
print(f"Всего строк данных: {len(df)}")
print(f"Всего столбцов: {len(df.columns)}")
print(f"Столбцы: {', '.join(df.columns)}")
